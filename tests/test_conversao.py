import zipfile
from datetime import datetime

import openpyxl
from lxml import etree

from conversor.conversao import MODELO_PADRAO, converter_pdf
from conversor.leitor_pdf import ler_previa
from conversor.mapeamento import Config
from conversor.planilha import Planilha

COLUNAS_VALORES = "DEFGHIJKLMNOPQR"


def gerar(previa, tmp_path, razao=None):
    config = Config()
    if razao:
        config.definir_razao_social(123, razao)
    (resultado,) = converter_pdf(previa, config, tmp_path / "saida")
    return resultado


def test_planilha_gerada(previa, tmp_path):
    resultado = gerar(previa, tmp_path, razao="EMPRESA EXEMPLO LTDA")
    assert resultado.destino.name == "dominio planilha - 0123 - 08-2026.xlsx"
    assert resultado.avisos == []
    ws = openpyxl.load_workbook(resultado.destino).active

    assert ws["C3"].value == 123
    assert ws["C4"].value == "EMPRESA EXEMPLO LTDA"
    assert ws["C5"].value == 12345678000190
    assert ws["C6"].value == datetime(2026, 8, 1)
    assert ws["C6"].number_format == "mm/yyyy"

    linhas = [[ws.cell(r, c).value for c in range(1, 19)] for r in range(11, 15)]
    #          A   B   C                    D        E     F     G      H      I    J     K     L     M     N      O    P     Q   R
    assert linhas[0] == [11, 5, "BELTRANA DE SOUZA", None, None, None, None, None, None, None, None, None, None, None, 275, None, None, 1.26]
    assert linhas[1] == [11, 10, "FULANO DE TAL", 1234.56, None, None, 450.1, 46.33, 100, None, None, None, None, 13.26, 300, None, 42, 32.76]
    assert linhas[2] == [11, 20, "CICRANO PEREIRA", 665.83, None, None, None, None, None, None, None, None, None, None, None, None, None, 15.12]
    assert all(v is None for v in linhas[3])

    assert ws["B77"].value == "=COUNTA(B11:B76)"
    assert ws["D77"].value == "=SUM(D11:D76)"


def test_totais_calculados_gravados(previa, tmp_path):
    resultado = gerar(previa, tmp_path)
    ws = openpyxl.load_workbook(resultado.destino, data_only=True).active
    assert ws["B77"].value == 3
    assert round(ws["D77"].value, 2) == 1900.39
    assert round(ws["R77"].value, 2) == 49.14
    assert ws["E77"].value == 0


def test_formatacao_do_modelo_preservada(previa, tmp_path):
    resultado = gerar(previa, tmp_path)
    modelo = openpyxl.load_workbook(MODELO_PADRAO).active
    gerado = openpyxl.load_workbook(resultado.destino).active
    for linha in list(range(1, 11)) + list(range(11, 78)):
        for coluna in range(1, 19):
            a, b = modelo.cell(linha, coluna), gerado.cell(linha, coluna)
            assert a._style == b._style, a.coordinate
            assert a.number_format == b.number_format, a.coordinate
            if linha <= 10 or linha == 77:
                assert a.value == b.value or linha in (3, 4, 5, 6), a.coordinate
    assert [c.width for c in modelo.column_dimensions.values()] == [
        c.width for c in gerado.column_dimensions.values()
    ]
    assert modelo.merged_cells.ranges == gerado.merged_cells.ranges


def test_partes_do_arquivo_mantidas(previa, tmp_path):
    resultado = gerar(previa, tmp_path)
    with zipfile.ZipFile(MODELO_PADRAO) as a, zipfile.ZipFile(resultado.destino) as b:
        assert a.namelist() == b.namelist()
        for nome in a.namelist():
            if nome not in ("xl/worksheets/sheet1.xml", "xl/sharedStrings.xml"):
                assert a.read(nome) == b.read(nome), nome
        sst = etree.fromstring(b.read("xl/sharedStrings.xml"))
        textos = ["".join(si.itertext()) for si in sst]
        assert len(textos) == len(set(textos)) == int(sst.get("uniqueCount"))


def test_mais_funcionarios_que_linhas_do_modelo(tmp_path):
    from decimal import Decimal

    from conversor.mapeamento import LinhaPlanilha

    planilha = Planilha(MODELO_PADRAO)
    capacidade = planilha.capacidade
    linhas = [LinhaPlanilha(11, i, f"PESSOA {i}", {"37": Decimal("10.5")}) for i in range(1, capacidade + 15)]
    planilha.preencher(1, "X", "123", None, linhas)
    destino = tmp_path / "grande.xlsx"
    planilha.salvar(destino)

    ws = openpyxl.load_workbook(destino).active
    total = 11 + len(linhas)
    assert ws.cell(total, 1).value == "TOTAL"
    assert ws.cell(total, 2).value == f"=COUNTA(B11:B{total - 1})"
    assert ws.cell(total, 4).value == f"=SUM(D11:D{total - 1})"
    assert ws.cell(total, 18).value == f"=SUM(R11:R{total - 1})"
    assert ws.cell(total - 1, 3).value == f"PESSOA {len(linhas)}"
    valores = openpyxl.load_workbook(destino, data_only=True).active
    assert valores.cell(total, 2).value == len(linhas)
    assert round(valores.cell(total, 4).value, 2) == round(10.5 * len(linhas), 2)
    with zipfile.ZipFile(destino) as z:
        assert "xl/calcChain.xml" not in z.namelist()


def test_varias_paginas(previa_longa, tmp_path):
    (resultado,) = converter_pdf(previa_longa, Config(), tmp_path)
    ws = openpyxl.load_workbook(resultado.destino).active
    assert [ws.cell(r, 2).value for r in range(11, 23)] == list(range(100, 112))
    assert ler_previa(previa_longa)[0].avisos() == []
